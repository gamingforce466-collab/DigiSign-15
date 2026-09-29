"""Ekspor laporan benchmark dan tamper ke workbook Excel."""
import io


def export_xlsx(report):
    """Rekap pengujian kuantitatif ke Excel (.xlsx), 3 sheet. Kembalikan bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F172A")
    title_font = Font(bold=True, size=13)
    ok_fill = PatternFill("solid", fgColor="DCFCE7")
    bad_fill = PatternFill("solid", fgColor="FEE2E2")
    center = Alignment(horizontal="center", vertical="center")

    def header(ws, values):
        ws.append(values)
        for cell in ws[ws.max_row]:
            cell.font, cell.fill, cell.alignment = head_font, head_fill, center

    def widths(ws, minimum=10, maximum=60, first_row=1):
        for col in ws.columns:
            longest = max((len(str(c.value)) for c in col
                           if c.value is not None and c.row >= first_row), default=0)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 2, minimum), maximum)

    n = report["trials"]
    rsa_r, ec_r = report["algorithms"]
    algos = report["algorithms"]
    wb = Workbook()

    ws1 = wb.active
    ws1.title = f"Benchmark {n}x"
    ws1.append([f"Benchmark {n} percobaan: waktu tanda tangan dan verifikasi (ms)"])
    ws1["A1"].font = title_font
    ws1.append([f"Dibuat: {report['created_at']}. Hash SHA-256 (32 byte) ditandatangani; satuan milidetik."])
    header(ws1, ["Percobaan"] + [f"{a['name']} - tanda tangan (ms)" if k == 0 else f"{a['name']} - verifikasi (ms)"
                                  for a in algos for k in (0, 1)])
    for i in range(n):
        row = [i + 1]
        for a in algos:
            row += [round(a["raw_sign"][i], 4), round(a["raw_verify"][i], 4)]
        ws1.append(row)
    ws1.freeze_panes = "B4"
    widths(ws1, 14, 46, first_row=3)
    ws1.column_dimensions["A"].width = 12

    ws2 = wb.create_sheet("Ringkasan Metrik")
    ws2.append(["Ringkasan metrik waktu dan ukuran kriptografi"])
    ws2["A1"].font = title_font
    ws2.append([])
    header(ws2, ["Algoritma", "Operasi", "Percobaan", "Mean (ms)", "Min (ms)", "Maks (ms)", "Std Dev (ms)"])
    for a in algos:
        for label, key in (("Tanda tangan", "sign"), ("Verifikasi", "verify")):
            st = a[key]
            ws2.append([a["name"], label, n, round(st["mean"], 4), round(st["min"], 4),
                        round(st["max"], 4), round(st["std"], 4)])
    ws2.append([])
    header(ws2, ["Algoritma", "Tanda tangan (byte)", "Kunci publik DER (byte)",
                 "Kunci publik PEM (byte)", "Ukuran hash (byte)"])
    for a in algos:
        z = a["sizes"]
        ws2.append([a["name"], z["signature_bytes"], z["public_key_der_bytes"],
                    z["public_key_pem_bytes"], 32])
    ws2.append([])
    ws2.append(["Perbandingan RSA-2048 terhadap ECDSA P-256"])
    ws2.cell(ws2.max_row, 1).font = Font(bold=True)
    rz, ez = rsa_r["sizes"], ec_r["sizes"]
    ws2.append(["Tanda tangan RSA / ECDSA", round(rz["signature_bytes"] / ez["signature_bytes"], 2), "kali lebih besar"])
    ws2.append(["Kunci publik DER RSA / ECDSA", round(rz["public_key_der_bytes"] / ez["public_key_der_bytes"], 2), "kali lebih besar"])
    ws2.append(["Tanda tangan RSA / ECDSA (waktu, mean)", round(rsa_r["sign"]["mean"] / ec_r["sign"]["mean"], 2), "kali lebih lama"])
    ws2.append(["Catatan: aplikasi memakai RSA-2048-PSS. ECDSA P-256 hanya pembanding. Ukuran ECDSA (DER) bisa 70-72 byte."])
    widths(ws2, 14, 40, first_row=3)
    ws2.column_dimensions["A"].width = 42

    t = report["tamper"]
    ws3 = wb.create_sheet("Uji Tamper 1-Byte")
    ws3.append(["Uji tamper: ubah 1 byte PDF bertanda tangan, verifikasi harus gagal (TAMPERED)"])
    ws3["A1"].font = title_font
    ws3.append([f"Ukuran PDF: {t['pdf_size']} byte. Hash asli SHA-256: {t['original_digest']}"])
    ws3.append([f"Kontrol (PDF asli tanpa perubahan): {'VALID' if t['control_valid'] else 'GAGAL'}"])
    header(ws3, ["No", "Offset byte", "Posisi (%)", "Byte asli", "Byte baru", "Hash setelah diubah (SHA-256)",
                 "Tanda tangan", "Status", "Diharapkan", "Hasil uji"])
    for r in t["rows"]:
        ws3.append([r["number"], r["offset"], r["position_pct"], r["old_byte"], r["new_byte"], r["new_digest"],
                    "Valid" if r["signature_valid"] else "Ditolak", r["status"], "TAMPERED",
                    "LULUS" if r["passed"] else "GAGAL"])
        fill = ok_fill if r["passed"] else bad_fill
        ws3.cell(ws3.max_row, 10).fill = fill
        ws3.cell(ws3.max_row, 8).fill = fill
    ws3.append([])
    passed = sum(1 for r in t["rows"] if r["passed"])
    ws3.append([f"Ringkasan: {passed} dari {len(t['rows'])} skenario berstatus TAMPERED."])
    ws3.cell(ws3.max_row, 1).font = Font(bold=True)
    ws3.freeze_panes = "A5"
    widths(ws3, 10, 70, first_row=4)

    for ws in (ws1, ws2, ws3):
        last_col = ws.max_column
        for row in list(ws.iter_rows()):
            filled = [c for c in row if c.value is not None]
            if len(filled) == 1 and filled[0].column == 1 and last_col > 1:
                r = filled[0].row
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
                if len(str(filled[0].value)) > 100:
                    ws.row_dimensions[r].height = 32
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()