"""Pengukuran waktu dan ukuran kriptografi.

RSA-2048-PSS adalah algoritma yang dipakai aplikasi. ECDSA P-256 hanya
pembanding ukuran dan kecepatan.
"""
import hashlib
import io
import os
import tempfile
import time
import statistics
from datetime import datetime

from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import ec, padding, rsa, utils

from crypto import signer as signer_module
from crypto import verifier as verifier_module

DEFAULT_TRIALS = 30


def _stats(values_ms):
    return {
        "mean": statistics.mean(values_ms),
        "min": min(values_ms),
        "max": max(values_ms),
        "std": statistics.stdev(values_ms) if len(values_ms) > 1 else 0.0,
    }


def _rsa_case():
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    return {
        "name": "RSA-2048-PSS (dipakai aplikasi)",
        "key": key,
        "sign": signer_module.sign_digest,
        "verify": verifier_module.verify_digest,
    }


def _ecdsa_case():
    key = ec.generate_private_key(ec.SECP256R1())
    prehashed = ec.ECDSA(utils.Prehashed(hashes.SHA256()))

    def sign(private_key, digest):
        return private_key.sign(digest, prehashed)

    def verify(public_key, digest, signature):
        try:
            public_key.verify(signature, digest, prehashed)
            return True
        except Exception:
            return False

    return {"name": "ECDSA P-256 (pembanding)", "key": key, "sign": sign, "verify": verify}


def _sizes(key, signature):
    public = key.public_key()
    return {
        "signature_bytes": len(signature),
        "public_key_der_bytes": len(public.public_bytes(
            serialization.Encoding.DER, serialization.PublicFormat.SubjectPublicKeyInfo)),
        "public_key_pem_bytes": len(public.public_bytes(
            serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo)),
    }


def run_benchmark(trials=DEFAULT_TRIALS):
    """Ukur waktu tanda tangan dan verifikasi sebanyak `trials` kali per algoritma."""
    if trials < 30:
        raise ValueError("Minimal 30 percobaan")
    digest = os.urandom(32)  # stand-in untuk hash SHA-256 dokumen
    results = []
    for case in (_rsa_case(), _ecdsa_case()):
        key, sign, verify = case["key"], case["sign"], case["verify"]
        public = key.public_key()
        signature = sign(key, digest)  # pemanasan, tidak dihitung
        verify(public, digest, signature)
        sign_ms, verify_ms = [], []
        for _ in range(trials):
            t0 = time.perf_counter()
            signature = sign(key, digest)
            sign_ms.append((time.perf_counter() - t0) * 1000)
            t0 = time.perf_counter()
            ok = verify(public, digest, signature)
            verify_ms.append((time.perf_counter() - t0) * 1000)
            if not ok:
                raise RuntimeError("Verifikasi gagal saat benchmark")
        results.append({
            "name": case["name"],
            "sign": _stats(sign_ms),
            "verify": _stats(verify_ms),
            "raw_sign": sign_ms,
            "raw_verify": verify_ms,
            "sizes": _sizes(key, signature),
        })
    return {"trials": trials, "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "algorithms": results}


TAMPER_FRACTIONS = (0.0, 0.05, 0.15, 0.25, 0.35, 0.50, 0.65, 0.75, 0.90, 1.0)


def run_tamper_tests(count=10):
    """Uji tamper: ubah 1 byte PDF bertanda tangan di `count` posisi berbeda.

    Setiap skenario harus berstatus TAMPERED (hash berubah dan tanda tangan ditolak).
    Alur sama dengan aplikasi: hash SHA-256 seluruh byte PDF, tanda tangan RSA-2048-PSS.
    """
    import randomdata
    from pdf import handler as pdf_module

    private_key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    public_key = private_key.public_key()
    pdf_bytes, _ = randomdata.make_random_pdf_bytes()
    size = len(pdf_bytes)

    def digest_of(data):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "uji.pdf")
            with open(path, "wb") as f:
                f.write(data)
            return pdf_module.compute_content_digest(path)

    original_digest = digest_of(pdf_bytes)
    signature = signer_module.sign_digest(private_key, original_digest)
    control_ok = verifier_module.verify_digest(public_key, original_digest, signature)

    offsets = []
    for fraction in TAMPER_FRACTIONS[:count]:
        offset = min(int(size * fraction), size - 1)
        while offset in offsets:
            offset += 1
        offsets.append(offset)

    rows = []
    for number, offset in enumerate(offsets, start=1):
        tampered = bytearray(pdf_bytes)
        old = tampered[offset]
        tampered[offset] = old ^ 0x01  # ubah tepat 1 byte
        new_digest = digest_of(bytes(tampered))
        sig_ok = verifier_module.verify_digest(public_key, new_digest, signature)
        status = "VALID" if (sig_ok and new_digest == original_digest) else "TAMPERED"
        rows.append({
            "number": number,
            "offset": offset,
            "position_pct": round(offset / size * 100, 1),
            "old_byte": f"0x{old:02X}",
            "new_byte": f"0x{tampered[offset]:02X}",
            "new_digest": new_digest.hex(),
            "signature_valid": sig_ok,
            "status": status,
            "passed": status == "TAMPERED",
        })
    return {
        "pdf_size": size,
        "original_digest": original_digest.hex(),
        "control_valid": control_ok,
        "rows": rows,
        "all_passed": control_ok and all(r["passed"] for r in rows),
    }


def run_full(trials=DEFAULT_TRIALS):
    """Benchmark waktu + ukuran + uji tamper 1-byte (10 skenario)."""
    report = run_benchmark(trials)
    report["tamper"] = run_tamper_tests(10)
    return report


def export_xlsx(report):
    """Rekap pengujian kuantitatif ke Excel (.xlsx), 3 sheet. Kembalikan bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F172A")
    title_font = Font(bold=True, size=13)
    ok_fill = PatternFill("solid", fgColor="DCFCE7")
    bad_fill = PatternFill("solid", fgColor="FEE2E2")
    center = Alignment(horizontal="center", vertical="center")

    def header(ws, row, values):
        ws.append(values)
        for cell in ws[ws.max_row]:
            cell.font, cell.fill, cell.alignment = head_font, head_fill, center

    def widths(ws, minimum=10, maximum=60, first_row=1):
        for col in ws.columns:
            longest = max((len(str(c.value)) for c in col
                           if c.value is not None and c.row >= first_row), default=0)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 2, minimum), maximum)

    n = report["trials"]
    rsa_r, ec_r = report["algorithms"]
    algos = report["algorithms"]
    wb = Workbook()

    # Sheet 1: benchmark per percobaan
    ws1 = wb.active
    ws1.title = f"Benchmark {n}x"
    ws1.append([f"Benchmark {n} percobaan: waktu tanda tangan dan verifikasi (ms)"])
    ws1["A1"].font = title_font
    ws1.append([f"Dibuat: {report['created_at']}. Hash SHA-256 (32 byte) ditandatangani; satuan milidetik."])
    header(ws1, 3, ["Percobaan"] + [f"{a['name']} - tanda tangan (ms)" if k == 0 else f"{a['name']} - verifikasi (ms)"
                                     for a in algos for k in (0, 1)])
    for i in range(n):
        row = [i + 1]
        for a in algos:
            row += [round(a["raw_sign"][i], 4), round(a["raw_verify"][i], 4)]
        ws1.append(row)
    ws1.freeze_panes = "B4"
    widths(ws1, 14, 46, first_row=3)
    ws1.column_dimensions["A"].width = 12

    # Sheet 2: ringkasan metrik dan ukuran
    ws2 = wb.create_sheet("Ringkasan Metrik")
    ws2.append(["Ringkasan metrik waktu dan ukuran kriptografi"])
    ws2["A1"].font = title_font
    ws2.append([])
    header(ws2, 3, ["Algoritma", "Operasi", "Percobaan", "Mean (ms)", "Min (ms)", "Maks (ms)", "Std Dev (ms)"])
    for a in algos:
        for label, key in (("Tanda tangan", "sign"), ("Verifikasi", "verify")):
            st = a[key]
            ws2.append([a["name"], label, n, round(st["mean"], 4), round(st["min"], 4),
                        round(st["max"], 4), round(st["std"], 4)])
    ws2.append([])
    header(ws2, 0, ["Algoritma", "Tanda tangan (byte)", "Kunci publik DER (byte)",
                                  "Kunci publik PEM (byte)", "Ukuran hash (byte)"])
    for a in algos:
        z = a["sizes"]
        ws2.append([a["name"], z["signature_bytes"], z["public_key_der_bytes"],
                    z["public_key_pem_bytes"], 32])
    ws2.append([])
    ws2.append(["Perbandingan RSA-2048 terhadap ECDSA P-256"])
    ws2.cell(ws2.max_row, 1).font = Font(bold=True)
    rz, ez = rsa_r["sizes"], ec_r["sizes"]
    ws2.append(["Tanda tangan RSA / ECDSA", round(rz["signature_bytes"] / ez["signature_bytes"], 2), "kali lebih besar"])
    ws2.append(["Kunci publik DER RSA / ECDSA", round(rz["public_key_der_bytes"] / ez["public_key_der_bytes"], 2), "kali lebih besar"])
    ws2.append(["Tanda tangan RSA / ECDSA (waktu, mean)", round(rsa_r["sign"]["mean"] / ec_r["sign"]["mean"], 2), "kali lebih lama"])
    ws2.append(["Catatan: aplikasi memakai RSA-2048-PSS. ECDSA P-256 hanya pembanding. Ukuran ECDSA (DER) bisa 70-72 byte."])
    widths(ws2, 14, 40, first_row=3)
    ws2.column_dimensions["A"].width = 42

    # Sheet 3: uji tamper 1 byte
    t = report["tamper"]
    ws3 = wb.create_sheet("Uji Tamper 1-Byte")
    ws3.append(["Uji tamper: ubah 1 byte PDF bertanda tangan, verifikasi harus gagal (TAMPERED)"])
    ws3["A1"].font = title_font
    ws3.append([f"Ukuran PDF: {t['pdf_size']} byte. Hash asli SHA-256: {t['original_digest']}"])
    ws3.append([f"Kontrol (PDF asli tanpa perubahan): {'VALID' if t['control_valid'] else 'GAGAL'}"])
    header(ws3, 4, ["No", "Offset byte", "Posisi (%)", "Byte asli", "Byte baru", "Hash setelah diubah (SHA-256)",
                    "Tanda tangan", "Status", "Diharapkan", "Hasil uji"])
    for r in t["rows"]:
        ws3.append([r["number"], r["offset"], r["position_pct"], r["old_byte"], r["new_byte"], r["new_digest"],
                    "Valid" if r["signature_valid"] else "Ditolak", r["status"], "TAMPERED",
                    "LULUS" if r["passed"] else "GAGAL"])
        fill = ok_fill if r["passed"] else bad_fill
        ws3.cell(ws3.max_row, 10).fill = fill
        ws3.cell(ws3.max_row, 8).fill = fill
    ws3.append([])
    passed = sum(1 for r in t["rows"] if r["passed"])
    ws3.append([f"Ringkasan: {passed} dari {len(t['rows'])} skenario berstatus TAMPERED."])
    ws3.cell(ws3.max_row, 1).font = Font(bold=True)
    ws3.freeze_panes = "A5"
    widths(ws3, 10, 70, first_row=4)

    # Rata tengah: semua sel. Baris judul/catatan (hanya kolom A terisi) digabung selebar tabel.
    for ws in (ws1, ws2, ws3):
        last_col = ws.max_column
        for row in list(ws.iter_rows()):
            filled = [c for c in row if c.value is not None]
            if len(filled) == 1 and filled[0].column == 1 and last_col > 1:
                r = filled[0].row
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
                if len(str(filled[0].value)) > 100:
                    ws.row_dimensions[r].height = 32
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    # Jalankan dari terminal: python -m crypto.benchmark [jumlah_percobaan]
    import sys
    n = int(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_TRIALS
    rep = run_benchmark(n)
    print(f"Benchmark {rep['trials']} percobaan (ms)")
    print(f"{'Algoritma':34}{'Operasi':14}{'Mean':>9}{'Min':>9}{'Maks':>9}{'Std':>9}")
    for a in rep["algorithms"]:
        for label, key in (("Tanda tangan", "sign"), ("Verifikasi", "verify")):
            st = a[key]
            print(f"{a['name']:34}{label:14}{st['mean']:9.4f}{st['min']:9.4f}{st['max']:9.4f}{st['std']:9.4f}")
    print()
    print(f"{'Algoritma':34}{'TTD (B)':>9}{'Pub DER (B)':>13}{'Pub PEM (B)':>13}")
    for a in rep["algorithms"]:
        z = a["sizes"]
        print(f"{a['name']:34}{z['signature_bytes']:9}{z['public_key_der_bytes']:13}{z['public_key_pem_bytes']:13}")
