import io
import os
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from reportlab.pdfgen import canvas
from streamlit.testing.v1 import AppTest

import app as app_module
import document_formats
from crypto import keys as key_module
from qr import generator as qr_module

ROOT = Path(__file__).resolve().parents[1]
STREAMLIT_APP = ROOT / "streamlit_app.py"
ACCESS_PASSWORD = "streamlit-demo-test"
OWNER_A = "streamlit_test_owner_a"
OWNER_B = "streamlit_test_owner_b"
PASS_A = "streamlit-test-pass-a"
PASS_B = "streamlit-test-pass-b"
MIME_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".txt": "text/plain",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


@pytest.fixture()
def demo_storage(tmp_path, monkeypatch):
    documents_dir = tmp_path / "documents"
    signatures_dir = tmp_path / "signatures"
    keys_dir = tmp_path / "keys"
    for folder in (documents_dir, signatures_dir, keys_dir):
        folder.mkdir()

    monkeypatch.setattr(app_module, "DOCUMENTS_DIR", documents_dir)
    monkeypatch.setattr(app_module, "SIGNATURES_DIR", signatures_dir)
    monkeypatch.setattr(key_module, "KEYS_DIR", keys_dir)
    monkeypatch.setattr(app_module, "VERIFY_BASE_URL", "http://localhost:8501")
    monkeypatch.setenv("STREAMLIT_ACCESS_PASSWORD", ACCESS_PASSWORD)
    monkeypatch.delenv("DIGISIGN_DATA_DIR", raising=False)
    monkeypatch.delenv("STREAMLIT_VERIFY_BASE_URL", raising=False)
    monkeypatch.delenv("ENABLE_TEST_TOOLS", raising=False)
    monkeypatch.delenv("ENABLE_DATA_RESET", raising=False)
    key_module.generate_keypair(OWNER_A, PASS_A)
    key_module.generate_keypair(OWNER_B, PASS_B)
    return documents_dir, signatures_dir, keys_dir


def _sample_document(extension):
    buffer = io.BytesIO()
    if extension == ".pdf":
        document = canvas.Canvas(buffer)
        document.drawString(70, 700, "Dokumen contoh asli untuk uji tanda tangan")
        document.save()
    elif extension == ".docx":
        document = Document()
        document.add_paragraph("Dokumen Word contoh untuk uji tanda tangan")
        document.save(buffer)
    elif extension == ".txt":
        return "Dokumen teks contoh untuk uji tanda tangan\n".encode("utf-8")
    elif extension == ".xlsx":
        workbook = Workbook()
        workbook.active["A1"] = "Dokumen Excel contoh untuk uji tanda tangan"
        workbook.save(buffer)
    else:
        image = Image.new("RGB", (800, 600), "white")
        image.save(buffer, format="JPEG" if extension in (".jpg", ".jpeg") else "PNG")
    return buffer.getvalue()


def _tamper_document(data, extension):
    if extension == ".pdf":
        return app_module.flip_one_byte(data)
    if extension == ".docx":
        document = Document(io.BytesIO(data))
        document.paragraphs[0].text += " diubah"
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
    if extension == ".txt":
        return data.replace(b"Dokumen teks contoh", b"Dokumen teks diubah", 1)
    if extension == ".xlsx":
        workbook = load_workbook(io.BytesIO(data))
        workbook[workbook.sheetnames[0]]["A1"] = "Dokumen Excel diubah"
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    with Image.open(io.BytesIO(data)) as source:
        image = source.convert("RGB")
    image.putpixel((0, 0), (255, 0, 0))
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG" if extension in (".jpg", ".jpeg") else "PNG")
    return buffer.getvalue()


def _logged_in_app():
    app = AppTest.from_file(str(STREAMLIT_APP), default_timeout=60).run()
    app.text_input[0].set_value(ACCESS_PASSWORD)
    app.button(key="FormSubmitter:access_form-Masuk").click().run()
    assert not app.exception, app.exception
    return app


def _verify_in_streamlit(data, filename, manual_doc_id="", public_key=None):
    app = _logged_in_app()
    app.radio(key="page").set_value("Verifikasi").run()
    extension = Path(filename).suffix.lower()
    app.file_uploader[0].upload(filename, data, MIME_TYPES[extension]).run(timeout=60)
    if manual_doc_id:
        app.text_input(key="verify_doc_id").set_value(manual_doc_id)
    if public_key is not None:
        app.file_uploader[1].upload("wrong_public.pem", public_key, "application/x-pem-file").run(timeout=60)
    app.button(key="FormSubmitter:verify_document_form-Verifikasi").click().run(timeout=60)
    assert not app.exception, app.exception
    return app


@pytest.mark.parametrize("extension", [".pdf", ".docx", ".jpg", ".jpeg", ".png", ".txt", ".xlsx"])
def test_streamlit_sign_and_verify_every_supported_format(demo_storage, extension):
    documents_dir, signatures_dir, _ = demo_storage
    app = _logged_in_app()
    app.radio(key="page").set_value("Tanda Tangan").run()
    app.text_input(key="signing_owner_id").set_value(OWNER_A).run()
    app.text_input(key="signing_passphrase").set_value(PASS_A)
    app.text_input(key="signer_name").set_value("Penandatangan Demo")
    sample = _sample_document(extension)
    filename = f"dokumen_asli{extension}"
    app.file_uploader[0].upload(filename, sample, MIME_TYPES[extension]).run(timeout=60)
    app.button(key="FormSubmitter:sign_document_form-Tandatangani").click().run(timeout=60)

    assert not app.exception, app.exception
    assert not app.error, [message.value for message in app.error]
    records = list(signatures_dir.glob("*.json"))
    assert len(records) == 1
    doc_id = records[0].stem
    signed_path = documents_dir / f"{doc_id}_signed{extension}"
    assert signed_path.is_file()
    decoded_qr = document_formats.extract_qr_metadata(signed_path, extension)
    assert decoded_qr and any(item.get("doc_id") == doc_id for item in decoded_qr), decoded_qr

    app.radio(key="page").set_value("Verifikasi").run()
    app.file_uploader[0].upload(signed_path.name, signed_path.read_bytes(), MIME_TYPES[extension]).run(timeout=60)
    app.button(key="FormSubmitter:verify_document_form-Verifikasi").click().run(timeout=60)
    assert not app.exception, app.exception
    assert any(message.value.startswith("VALID:") for message in app.success), app.session_state.get("last_verification")
    assert any(button.label == "Unduh dokumen terverifikasi" for button in app.download_button)
    tampered = _tamper_document(signed_path.read_bytes(), extension)
    tampered_result = _verify_in_streamlit(tampered, signed_path.name, doc_id)
    assert any(message.value.startswith("TAMPERED:") for message in tampered_result.error)
    assert not any(button.label == "Unduh dokumen terverifikasi" for button in tampered_result.download_button)


@pytest.mark.parametrize(
    ("scenario", "expected"),
    [
        ("valid", "VALID"),
        ("tampered", "TAMPERED"),
        ("wrong_key", "KEY_MISMATCH"),
        ("fake_qr", "QR_FORGED"),
        ("unknown_doc", "NOT_FOUND"),
        ("no_qr", "NO_QR"),
        ("pasted_qr", "TAMPERED"),
    ],
)
def test_streamlit_verification_scenarios(demo_storage, scenario, expected):
    documents_dir, _, keys_dir = demo_storage
    original = _sample_document(".pdf")
    doc_id, _ = app_module.perform_signing(
        OWNER_A, PASS_A, "Penandatangan Demo", "Ketua", "Universitas Demo",
        "2026-09-30", pdf_bytes=original, original_filename="dokumen_asli.pdf",
    )
    signed_path = documents_dir / f"{doc_id}_signed.pdf"
    signed = signed_path.read_bytes()
    manual_doc_id = ""
    public_key = None

    if scenario == "tampered":
        signed = app_module.flip_one_byte(signed)
        manual_doc_id = doc_id
    elif scenario == "wrong_key":
        manual_doc_id = doc_id
        public_key = (keys_dir / f"{OWNER_B}_public.pem").read_bytes()
    elif scenario == "fake_qr":
        signed, _ = app_module.make_fake_qr_pdf()
    elif scenario == "unknown_doc":
        signed = original
        manual_doc_id = "abcdef123456"
    elif scenario == "no_qr":
        signed = original
    elif scenario == "pasted_qr":
        signed = app_module.paste_real_qr_on_other_pdf(signed)

    app = _verify_in_streamlit(signed, "hasil_uji.pdf", manual_doc_id, public_key)
    messages = [message.value for message in app.success] + [message.value for message in app.error]
    assert any(message.startswith(f"{expected}:") for message in messages), messages


def test_streamlit_supports_multiple_signers_and_old_versions(demo_storage):
    documents_dir, signatures_dir, _ = demo_storage
    app = _logged_in_app()
    app.radio(key="page").set_value("Tanda Tangan").run()
    app.text_input(key="signing_owner_id").set_value(OWNER_A).run()
    app.text_input(key="signing_passphrase").set_value(PASS_A)
    app.text_input(key="signer_name").set_value("Penandatangan Satu")
    app.file_uploader[0].upload("dokumen_asli.pdf", _sample_document(".pdf"), MIME_TYPES[".pdf"]).run(timeout=60)
    app.button(key="FormSubmitter:sign_document_form-Tandatangani").click().run(timeout=60)
    assert not app.exception, app.exception

    doc_id = next(signatures_dir.glob("*.json")).stem
    version_one = (documents_dir / f"{doc_id}_signed.pdf").read_bytes()
    app.radio(key="page").set_value("Tanda Tangan").run()
    app.selectbox[0].set_value(doc_id).run()
    app.text_input(key="signing_owner_id").set_value(OWNER_B).run()
    app.text_input(key="signing_passphrase").set_value(PASS_B)
    app.text_input(key="signer_name").set_value("Penandatangan Dua")
    app.button(key="FormSubmitter:sign_document_form-Tandatangani").click().run(timeout=60)
    assert not app.exception, app.exception
    version_two = (documents_dir / f"{doc_id}_signed.pdf").read_bytes()
    record = app_module.load_signature_record(doc_id)
    assert len(record["signers"]) == 2
    assert version_one != version_two

    for version, expected_in_document in ((version_one, 1), (version_two, 2)):
        result = _verify_in_streamlit(version, "signed.pdf", doc_id)
        assert any(message.value.startswith("VALID:") for message in result.success)
        metrics = {metric.label: metric.value for metric in result.metric}
        assert metrics["QR/metadata ditemukan"] == str(expected_in_document)
        assert metrics["Total penandatangan"] == "2"


def test_streamlit_add_signer_button_navigates_to_selected_document(demo_storage):
    documents_dir, signatures_dir, _ = demo_storage
    doc_id, _ = app_module.perform_signing(
        OWNER_A, PASS_A, "Penandatangan Satu", "Ketua", "Universitas Demo",
        "2026-09-30", pdf_bytes=_sample_document(".xlsx"), original_filename="Sheet 1-2.xlsx",
    )
    assert (documents_dir / f"{doc_id}_signed.xlsx").is_file()
    assert (signatures_dir / f"{doc_id}.json").is_file()

    app = _logged_in_app()
    app.button(key=f"add_signer_{doc_id}").click().run()
    assert not app.exception, app.exception
    assert app.radio(key="page").value == "Tanda Tangan"
    assert app.selectbox[0].value == doc_id
    app.text_input(key="signing_owner_id").set_value(OWNER_B).run()
    app.text_input(key="signing_passphrase").set_value(PASS_B)
    app.text_input(key="signer_name").set_value("Penandatangan Dua")
    app.button(key="FormSubmitter:sign_document_form-Tandatangani").click().run(timeout=60)
    assert not app.exception, app.exception
    assert not app.error, [message.value for message in app.error]
    record = app_module.load_signature_record(doc_id)
    assert len(record["signers"]) == 2
    signed_path = documents_dir / f"{doc_id}_signed.xlsx"
    assert signed_path.is_file()
    verification = _verify_in_streamlit(signed_path.read_bytes(), signed_path.name)
    assert any(message.value.startswith("VALID:") for message in verification.success)


def test_successful_signing_shows_actions_clears_form_and_verifies_stored_file(demo_storage):
    app = _logged_in_app()
    app.radio(key="page").set_value("Tanda Tangan").run()
    app.text_input(key="signing_owner_id").set_value(OWNER_A).run()
    app.text_input(key="signing_passphrase").set_value(PASS_A)
    app.text_input(key="signer_name").set_value("Penandatangan UI")
    app.file_uploader[0].upload("alur-mulus.pdf", _sample_document(".pdf"), MIME_TYPES[".pdf"]).run()
    app.button(key="FormSubmitter:sign_document_form-Tandatangani").click().run(timeout=60)

    assert not app.exception, app.exception
    doc_id = app.session_state["last_signed_doc_id"]
    assert app.text_input(key="signing_passphrase").value == ""
    assert app.text_input(key="signer_name").value == ""
    assert app.file_uploader[0].value is None
    assert any(button.label == "Unduh dokumen bertanda tangan" for button in app.download_button)
    assert any(button.label == "Verifikasi hasil tersimpan" for button in app.button)
    assert any(button.label == "Tambah penandatangan" for button in app.button)

    app.button(key=f"verify_result_{doc_id}").click().run(timeout=60)
    assert not app.exception, app.exception
    assert app.radio(key="page").value == "Verifikasi"
    assert app.session_state["last_verification"]["status"] == "VALID"


def test_streamlit_demo_controls_and_benchmark_render(demo_storage):
    app = _logged_in_app()
    assert any(button.label == "Bersihkan semua data" for button in app.button)
    app.radio(key="page").set_value("Tanda Tangan").run()
    assert any(button.label == "Isi data kunci acak" for button in app.button)
    assert any(button.label == "Isi data penandatangan acak" for button in app.button)
    app.radio(key="page").set_value("Uji Kuantitatif").run()
    app.button(key="FormSubmitter:benchmark_form-Jalankan pengujian").click().run(timeout=90)
    assert not app.exception, app.exception
    assert any("LULUS" in metric.value for metric in app.metric)
    assert app.session_state["benchmark_report"]["tamper"]["all_passed"] is True


def test_manual_key_and_document_persist_without_demo_page(demo_storage):
    documents_dir, signatures_dir, keys_dir = demo_storage
    app = _logged_in_app()
    app.radio(key="page").set_value("Tanda Tangan").run()
    app.text_input(key="signing_owner_id").set_value(OWNER_A).run()
    app.text_input(key="signing_passphrase").set_value(PASS_A)
    app.text_input(key="signer_name").set_value("Penandatangan Manual")
    app.file_uploader[0].upload("manual-retained.txt", _sample_document(".txt"), MIME_TYPES[".txt"]).run()
    app.button(key="FormSubmitter:sign_document_form-Tandatangani").click().run(timeout=60)
    assert not app.exception, app.exception

    record_path = next(signatures_dir.glob("*.json"))
    doc_id = record_path.stem
    assert key_module.key_exists(OWNER_A)
    assert (keys_dir / f"{OWNER_A}_private.pem").is_file()
    assert (documents_dir / f"{doc_id}_original.txt").is_file()
    assert (documents_dir / f"{doc_id}_signed.txt").is_file()

    reopened = _logged_in_app()
    assert "Demo End-to-End" not in reopened.radio(key="page").options
    assert next(metric for metric in reopened.metric if metric.label == "Dokumen").value == "1"
    assert any("manual-retained.txt" in item.value for item in reopened.markdown)


def test_streamlit_random_key_and_signer_controls_work(demo_storage):
    _, _, keys_dir = demo_storage
    app = _logged_in_app()
    app.radio(key="page").set_value("Tanda Tangan").run()
    next(button for button in app.button if button.label == "Isi data kunci acak").click().run()
    owner_id = app.text_input(key="new_key_owner").value
    passphrase = app.text_input(key="new_key_passphrase").value
    assert owner_id and passphrase
    app.button(key="FormSubmitter:create_key_form-Buat kunci").click().run()
    assert key_module.key_exists(owner_id)
    assert (keys_dir / f"{owner_id}_private.pem").is_file()
    assert app.text_input(key="signing_owner_id").value == owner_id
    assert any(button.label == "Unduh private key terenkripsi" for button in app.download_button)
    assert any(button.label == "Unduh public key" for button in app.download_button)

    next(button for button in app.button if button.label == "Isi data penandatangan acak").click().run()
    assert app.text_input(key="signer_name").value
    assert app.text_input(key="signer_position").value
    assert app.text_input(key="signer_institution").value


def test_signing_key_name_survives_ten_streamlit_reruns(demo_storage):
    app = _logged_in_app()
    app.radio(key="page").set_value("Tanda Tangan").run()
    for _ in range(10):
        app.text_input(key="signing_owner_id").set_value(OWNER_A).run()
        assert app.text_input(key="signing_owner_id").value == OWNER_A


def test_streamlit_clear_data_action_works_after_confirmation(demo_storage):
    documents_dir, signatures_dir, keys_dir = demo_storage
    app = _logged_in_app()
    app.checkbox[0].check().run()
    app.button[0].click().run()
    assert not app.exception, app.exception
    assert list(documents_dir.iterdir()) == []
    assert list(signatures_dir.iterdir()) == []
    assert list(keys_dir.iterdir()) == []


def test_streamlit_login_rejects_wrong_password_and_logout_returns_to_login(demo_storage):
    app = AppTest.from_file(str(STREAMLIT_APP), default_timeout=60).run()
    app.text_input[0].set_value("password-salah")
    app.button(key="FormSubmitter:access_form-Masuk").click().run()
    assert any("Password akses salah" in message.value for message in app.error)
    assert not app.radio

    app.text_input[0].set_value(ACCESS_PASSWORD)
    app.button(key="FormSubmitter:access_form-Masuk").click().run()
    assert not app.exception, app.exception
    assert app.radio(key="page").value == "Beranda"
    next(button for button in app.button if button.label == "Keluar").click().run()
    assert not app.exception, app.exception
    assert app.text_input[0].label == "Password akses"
    assert not app.radio


def test_streamlit_deep_link_opens_verification_with_doc_id(demo_storage):
    app = AppTest.from_file(str(STREAMLIT_APP), default_timeout=60)
    app.query_params["doc_id"] = "abcdef123456"
    app.run()
    app.text_input[0].set_value(ACCESS_PASSWORD)
    app.button(key="FormSubmitter:access_form-Masuk").click().run()
    assert not app.exception, app.exception
    assert app.radio(key="page").value == "Verifikasi"
    assert app.text_input(key="verify_doc_id").value == "abcdef123456"
    assert app.session_state["last_verification"]["status"] == "NOT_FOUND"


def test_streamlit_qr_deep_link_verifies_stored_document_and_checks_metadata(demo_storage):
    doc_id, record = app_module.perform_signing(
        OWNER_A, PASS_A, "Penandatangan QR", "Ketua", "Universitas Demo",
        "2026-09-30", pdf_bytes=_sample_document(".pdf"), original_filename="qr-demo.pdf",
    )
    signer = record["signers"][0]
    qr_metadata = {
        "doc_id": doc_id,
        "signer_id": signer["signer_id"],
        "name": signer["signer_name"],
        "position": signer["position"],
        "institution": signer["institution"],
        "date": signer["signed_date"],
        "hash": record["content_hash_hex"],
        "verify_url": f"http://localhost:8501?doc_id={doc_id}",
    }
    decoded_qr = qr_module.decode_qr_image(qr_module.generate_qr_image(qr_metadata))
    query = parse_qs(urlsplit(decoded_qr["verify_url"]).query)
    app = AppTest.from_file(str(STREAMLIT_APP), default_timeout=60)
    app.query_params["doc_id"] = doc_id
    app.query_params["m"] = query["m"][0]
    app.run()

    assert not app.exception, app.exception
    assert not app.radio
    assert not app.text_input
    assert any(message.value.startswith("VALID:") for message in app.success)
    assert any(button.label == "Unduh dokumen terverifikasi" for button in app.download_button)
    assert app.session_state["last_verification"]["status"] == "VALID"

    qr_metadata["name"] = "Pemalsu"
    forged_qr = qr_module.decode_qr_image(qr_module.generate_qr_image(qr_metadata))
    app.query_params["m"] = parse_qs(urlsplit(forged_qr["verify_url"]).query)["m"][0]
    app.run()
    assert not app.exception, app.exception
    assert not app.radio
    assert app.session_state["last_verification"]["status"] == "QR_FORGED"
    assert not any(button.label == "Unduh dokumen terverifikasi" for button in app.download_button)


def test_streamlit_sign_form_reports_invalid_password_and_invalid_document(demo_storage):
    _, signatures_dir, _ = demo_storage
    app = _logged_in_app()
    app.radio(key="page").set_value("Tanda Tangan").run()
    app.text_input(key="signing_owner_id").set_value(OWNER_A).run()
    app.text_input(key="signing_passphrase").set_value("incorrect-passphrase")
    app.text_input(key="signer_name").set_value("Penandatangan Uji")
    app.file_uploader[0].upload("dokumen.pdf", _sample_document(".pdf"), MIME_TYPES[".pdf"]).run()
    app.button(key="FormSubmitter:sign_document_form-Tandatangani").click().run(timeout=60)
    assert any("Passphrase salah" in message.value for message in app.error)
    assert list(signatures_dir.glob("*.json")) == []

    app = _logged_in_app()
    app.radio(key="page").set_value("Tanda Tangan").run()
    app.text_input(key="signing_owner_id").set_value(OWNER_A).run()
    app.text_input(key="signing_passphrase").set_value(PASS_A)
    app.text_input(key="signer_name").set_value("Penandatangan Uji")
    app.file_uploader[0].upload("rusak.pdf", b"bukan berkas PDF", MIME_TYPES[".pdf"]).run()
    app.button(key="FormSubmitter:sign_document_form-Tandatangani").click().run(timeout=60)
    assert any("bukan PDF yang valid" in message.value for message in app.error)
    assert list(signatures_dir.glob("*.json")) == []


def test_streamlit_verification_reports_missing_file_and_invalid_public_key(demo_storage):
    documents_dir, _, _ = demo_storage
    doc_id, _ = app_module.perform_signing(
        OWNER_A, PASS_A, "Penandatangan Uji", pdf_bytes=_sample_document(".pdf"),
        original_filename="dokumen.pdf",
    )
    signed_path = documents_dir / f"{doc_id}_signed.pdf"

    app = _logged_in_app()
    app.radio(key="page").set_value("Verifikasi").run()
    app.button(key="FormSubmitter:verify_document_form-Verifikasi").click().run()
    assert any("Pilih dokumen" in message.value for message in app.error)

    app = _logged_in_app()
    app.radio(key="page").set_value("Verifikasi").run()
    app.file_uploader[0].upload(signed_path.name, signed_path.read_bytes(), MIME_TYPES[".pdf"]).run()
    app.file_uploader[1].upload("invalid.pem", b"not a PEM key", "application/x-pem-file").run()
    app.button(key="FormSubmitter:verify_document_form-Verifikasi").click().run(timeout=60)
    assert any("Kunci publik tidak valid" in message.value for message in app.error)