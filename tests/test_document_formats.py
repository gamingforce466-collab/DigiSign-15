import io

import pytest
from PIL import Image
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

import document_formats
from qr import generator as qr_module


def sample_bytes(extension):
    if extension == ".pdf":
        buffer = io.BytesIO()
        pdf = canvas.Canvas(buffer)
        pdf.drawString(70, 700, "Document format test")
        pdf.save()
        return buffer.getvalue()
    if extension == ".docx":
        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("Document format test")
        heading_style = document.styles["Heading 1"]
        heading_style._element.getparent().remove(heading_style._element)
        document.save(buffer)
        return buffer.getvalue()
    if extension == ".txt":
        return "Document format test\n".encode("utf-8")
    if extension == ".xlsx":
        buffer = io.BytesIO()
        workbook = Workbook()
        workbook.active["A1"] = "Document format test"
        workbook.save(buffer)
        return buffer.getvalue()

    buffer = io.BytesIO()
    Image.new("RGB", (800, 600), "white").save(
        buffer, format="JPEG" if extension in (".jpg", ".jpeg") else "PNG"
    )
    return buffer.getvalue()


@pytest.mark.parametrize("extension", [".pdf", ".docx", ".jpg", ".jpeg", ".png", ".txt", ".xlsx"])
def test_supported_documents_keep_format_and_qr_metadata(tmp_path, extension):
    source_path = tmp_path / f"source{extension}"
    signed_path = tmp_path / f"signed{extension}"
    source_path.write_bytes(sample_bytes(extension))
    qr_metadata = [
        {
            "doc_id": "a1b2c3d4e5f6",
            "signer_id": f"signer-{index}",
            "name": f"Signer {index}",
            "position": "Ketua",
            "institution": "Universitas Contoh Nusantara",
            "date": "2026-09-29",
            "hash": "a" * 64,
            "verify_url": f"http://localhost:5000/verify?doc_id=a1b2c3d4e5f6",
        }
        for index in (1, 2)
    ]
    qr_images = [qr_module.generate_qr_image(metadata) for metadata in qr_metadata]
    signers = [{"name": item["name"], "position": "Ketua", "institution": "Univ", "date": "2026-09-29"}
               for item in qr_metadata]

    assert document_formats.is_valid_document(source_path.read_bytes(), source_path.name)
    document_formats.add_signature_marks(source_path, signed_path, extension, qr_images, signers)

    assert document_formats.is_valid_document(signed_path.read_bytes(), signed_path.name)
    decoded = document_formats.extract_qr_metadata(signed_path, extension)
    assert {item["signer_id"] for item in decoded} == {"signer-1", "signer-2"}


def test_format_and_file_contents_must_match():
    png_data = sample_bytes(".png")
    assert document_formats.is_valid_document(png_data, "image.png")
    assert not document_formats.is_valid_document(png_data, "image.jpg")
    assert not document_formats.is_valid_document(b"not a document", "file.docx")
    assert not document_formats.is_valid_document(b"not a workbook", "file.xlsx")