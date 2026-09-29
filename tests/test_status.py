"""Tes status verifikasi: VALID, TAMPERED, KEY_MISMATCH, QR_FORGED, dan 3 penandatangan."""
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pytest

import app as app_module
import randomdata
from crypto import keys as key_module

OWNERS = [("pytest_status_a", "pass_status_a"), ("pytest_status_b", "pass_status_b"),
          ("pytest_status_c", "pass_status_c")]


@pytest.fixture()
def env(tmp_path, monkeypatch):
    docs = tmp_path / "documents"
    sigs = tmp_path / "signatures"
    docs.mkdir()
    sigs.mkdir()
    monkeypatch.setattr(app_module, "DOCUMENTS_DIR", docs)
    monkeypatch.setattr(app_module, "SIGNATURES_DIR", sigs)
    for owner, pw in OWNERS:
        if not key_module.key_exists(owner):
            key_module.generate_keypair(owner, pw)
    return docs, sigs


def sign_three(docs):
    pdf_bytes, _ = randomdata.make_random_pdf_bytes()
    versions = []
    doc_id = ""
    for n, (owner, pw) in enumerate(OWNERS, start=1):
        doc_id, _ = app_module.perform_signing(
            owner, pw, f"Penandatangan {n}", "Ketua", "Univ Uji", "2026-09-28",
            doc_id=doc_id, pdf_bytes=pdf_bytes if n == 1 else None, original_filename="surat.pdf")
        versions.append((docs / f"{doc_id}_signed.pdf").read_bytes())
    return doc_id, versions


def test_valid_status(env):
    docs, _ = env
    doc_id, versions = sign_three(docs)
    res = app_module.verify_pdf_bytes(versions[-1])
    assert res["status"] == "VALID" and res["overall_valid"] is True


def test_tampered_status(env):
    docs, _ = env
    doc_id, versions = sign_three(docs)
    res = app_module.verify_pdf_bytes(app_module.flip_one_byte(versions[-1]), doc_id)
    assert res["status"] == "TAMPERED" and res["overall_valid"] is False


def test_key_mismatch_status(env):
    docs, _ = env
    doc_id, versions = sign_three(docs)
    res = app_module.verify_pdf_bytes(versions[-1], doc_id, key_module.generate_ephemeral_public_key())
    assert res["status"] == "KEY_MISMATCH"
    assert res["integrity_ok"] is True
    assert not any(s["signature_valid"] for s in res["signers"])


def test_fabricated_qr_is_forged(env):
    pdf, fake = app_module.make_fake_qr_pdf()
    res = app_module.verify_pdf_bytes(pdf)
    assert res["status"] == "QR_FORGED"
    assert res["doc_id"] == fake["doc_id"]


def test_real_qr_pasted_on_other_pdf_is_rejected(env):
    docs, _ = env
    doc_id, versions = sign_three(docs)
    res = app_module.verify_pdf_bytes(app_module.paste_real_qr_on_other_pdf(versions[-1]))
    assert res["status"] == "TAMPERED" and res["overall_valid"] is False


def test_three_signers_and_older_versions(env):
    docs, _ = env
    doc_id, versions = sign_three(docs)
    counts = []
    for v in versions:
        res = app_module.verify_pdf_bytes(v, doc_id)
        assert res["integrity_ok"] is True
        counts.append((res["signers_in_document"], res["signers_total"]))
    assert counts == [(1, 3), (2, 3), (3, 3)]


def test_no_qr_and_no_doc_id(env):
    pdf, _ = randomdata.make_random_pdf_bytes()
    assert app_module.verify_pdf_bytes(pdf)["status"] == "NO_QR"


def test_unknown_manual_doc_id_is_not_found(env):
    pdf, _ = randomdata.make_random_pdf_bytes()
    assert app_module.verify_pdf_bytes(pdf, "abcdef123456")["status"] == "NOT_FOUND"


def test_invalid_owner_name_rejected(env):
    with pytest.raises(app_module.SigningError):
        app_module.perform_signing("../evil", "x", "Nama", pdf_bytes=b"x")


def test_benchmark_page_and_excel_download(env):
    import io
    from openpyxl import load_workbook
    app_module.app.config["TESTING"] = True
    c = app_module.app.test_client()
    assert c.get("/benchmark").status_code == 200
    html = c.post("/benchmark", data={"trials": "30"}).get_data(as_text=True)
    assert "Unduh Rekapitulasi Data Pengujian (.xlsx)" in html
    assert "SEMUA LULUS" in html and "10 skenario" in html
    resp = c.get("/benchmark/xlsx")
    assert resp.status_code == 200 and resp.data[:2] == b"PK"
    assert "spreadsheetml" in resp.mimetype
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Benchmark 30x", "Ringkasan Metrik", "Uji Tamper 1-Byte"]
    s1, s2, s3 = wb.worksheets
    # sheet 1: 30 baris percobaan, semua angka positif
    rows = [r for r in s1.iter_rows(min_row=4, values_only=True) if r[0] is not None]
    assert len(rows) == 30 and all(v > 0 for r in rows for v in r[1:])
    # sheet 2: ukuran RSA
    text = [str(v) for r in s2.iter_rows(values_only=True) for v in r if v is not None]
    assert "256" in text and "294" in text
    # sheet 3: 10 skenario, semuanya TAMPERED dan LULUS
    tamper = [r for r in s3.iter_rows(min_row=5, values_only=True) if isinstance(r[0], int)]
    assert len(tamper) == 10
    assert all(r[7] == "TAMPERED" and r[9] == "LULUS" for r in tamper)
    assert len({r[1] for r in tamper}) == 10  # posisi byte berbeda-beda
    # semua sel berisi rata tengah
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    assert cell.alignment.horizontal == "center", (ws.title, cell.coordinate)


def test_excel_download_works_without_running_first(env):
    app_module.app.config["TESTING"] = True
    app_module.app.config["LAST_QUANT"] = None
    resp = app_module.app.test_client().get("/benchmark/xlsx")
    assert resp.status_code == 200 and resp.data[:2] == b"PK"


def test_web_autotest_is_removed(env):
    app_module.app.config["TESTING"] = True
    c = app_module.app.test_client()
    assert c.get("/autotest").status_code == 404
    assert "Uji Otomatis" not in c.get("/").get_data(as_text=True)


def test_clear_removes_everything_but_gitkeep(env, tmp_path, monkeypatch):
    docs, sigs = env
    keys = tmp_path / "keys"
    keys.mkdir()
    (keys / "x_private.pem").write_text("k")
    (docs / ".gitkeep").write_text("")
    doc_id, _ = sign_three(docs)
    monkeypatch.setattr(key_module, "KEYS_DIR", keys)
    assert list(sigs.glob("*.json")) and (docs / f"{doc_id}_signed.pdf").exists()
    app_module.app.config["TESTING"] = True
    c = app_module.app.test_client()
    resp = c.post("/clear", follow_redirects=True)
    assert "Data dibersihkan" in resp.get_data(as_text=True)
    assert list(sigs.glob("*.json")) == []
    assert [f.name for f in docs.iterdir()] == [".gitkeep"]
    assert list(keys.iterdir()) == []
    assert "Belum ada dokumen" in c.get("/").get_data(as_text=True)


def test_manual_pages_have_no_test_files_or_preview(env):
    app_module.app.config["TESTING"] = True
    c = app_module.app.test_client()
    sign_html = c.get("/sign").get_data(as_text=True)
    assert "use_sample_pdf" not in sign_html and "PDF contoh" not in sign_html
    assert "btn-random-sign" in sign_html  # Isi acak tetap ada
    assert "<iframe" not in sign_html and "pdf-preview" not in sign_html
    verify_html = c.get("/verify").get_data(as_text=True)
    assert "Uji cepat" not in verify_html and "scenario" not in verify_html


def test_verify_page_opens_and_is_named_verifikasi(env):
    app_module.app.config["TESTING"] = True
    c = app_module.app.test_client()
    resp = c.get("/verify")
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert "Verifikasi Manual" not in html and "autotest" not in html
    assert ">Verifikasi</a>" in html
    assert c.get("/verify?doc_id=abcdef123456").status_code == 200


def test_verify_upload_via_web_all_statuses(env):
    import io
    docs, _ = env
    doc_id, versions = sign_three(docs)
    app_module.app.config["TESTING"] = True
    c = app_module.app.test_client()

    def post(data_bytes, **extra):
        form = {"document": (io.BytesIO(data_bytes), "x.pdf"), **extra}
        return c.post("/verify", data=form, content_type="multipart/form-data").get_data(as_text=True)

    ok = post(versions[-1])
    assert '<p class="text-2xl font-bold">VALID</p>' in ok and "<b>3 dari 3</b>" in ok
    tam = post(app_module.flip_one_byte(versions[-1]), doc_id=doc_id)
    assert '<p class="text-2xl font-bold">TAMPERED</p>' in tam
    fake_pdf, _ = app_module.make_fake_qr_pdf()
    assert '<p class="text-2xl font-bold">QR_FORGED</p>' in post(fake_pdf)
    assert c.post("/verify", data={}, follow_redirects=False).status_code == 302
