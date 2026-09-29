import io
import sys
import os
import json
import hashlib

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from reportlab.pdfgen import canvas

import app as app_module
from crypto import keys as key_module

OWNER_A = "pytest_flow_owner_a"
OWNER_B = "pytest_flow_owner_b"
PASS_A = "pytest_flow_pass_a"
PASS_B = "pytest_flow_pass_b"


def status_block(code):
    return f'<p class="text-2xl font-bold">{code}</p>'


STATUS_VALID = status_block("VALID")


def make_pdf_bytes(text):
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, text)
    c.save()
    return buf.getvalue()


@pytest.fixture()
def client(tmp_path, monkeypatch):
    docs = tmp_path / "documents"
    sigs = tmp_path / "signatures"
    docs.mkdir()
    sigs.mkdir()
    monkeypatch.setattr(app_module, "DOCUMENTS_DIR", docs)
    monkeypatch.setattr(app_module, "SIGNATURES_DIR", sigs)
    for owner, pw in ((OWNER_A, PASS_A), (OWNER_B, PASS_B)):
        if not key_module.key_exists(owner):
            key_module.generate_keypair(owner, pw)
    app_module.app.config["TESTING"] = True
    return app_module.app.test_client(), docs, sigs


def sign(client, owner, passphrase, name, pdf_bytes=None, doc_id="", filename="surat.pdf"):
    data = {
        "doc_id": doc_id,
        "owner_id": owner,
        "passphrase": passphrase,
        "signer_name": name,
        "position": "Ketua",
        "institution": "Univ Uji",
    }
    if pdf_bytes is not None:
        data["document"] = (io.BytesIO(pdf_bytes), filename)
    return client.post("/sign", data=data, content_type="multipart/form-data")


def verify(client, pdf_bytes, doc_id, filename="cek.pdf"):
    resp = client.post(
        "/verify",
        data={"doc_id": doc_id, "document": (io.BytesIO(pdf_bytes), filename)},
        content_type="multipart/form-data",
    )
    return resp.get_data(as_text=True)


def only_doc_id(sigs):
    files = list(sigs.glob("*.json"))
    assert len(files) == 1
    return files[0].stem


def make_document_bytes(extension):
    buffer = io.BytesIO()
    if extension == ".docx":
        document = Document()
        document.add_paragraph("Lembar pengesahan uji format")
        document.save(buffer)
    elif extension == ".txt":
        buffer.write(b"Lembar pengesahan uji format\n")
    elif extension == ".xlsx":
        workbook = Workbook()
        workbook.active["A1"] = "Lembar pengesahan uji format"
        workbook.save(buffer)
    else:
        Image.new("RGB", (800, 600), "white").save(
            buffer, format="JPEG" if extension in (".jpg", ".jpeg") else "PNG"
        )
    return buffer.getvalue()


def tamper_document(data, extension):
    if extension == ".docx":
        document = Document(io.BytesIO(data))
        document.paragraphs[0].text += " diubah"
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
    if extension == ".txt":
        return data.replace(b"uji format", b"uji diubah", 1)
    if extension == ".xlsx":
        workbook = load_workbook(io.BytesIO(data))
        workbook[workbook.sheetnames[0]]["A1"] = "Lembar pengesahan diubah"
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    with Image.open(io.BytesIO(data)) as opened:
        image = opened.convert("RGB")
    image.putpixel((0, 0), (255, 0, 0))
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG" if extension in (".jpg", ".jpeg") else "PNG")
    return buffer.getvalue()


def test_sign_then_verify_signed_pdf_passes(client):
    c, docs, sigs = client
    sign(c, OWNER_A, PASS_A, "Budi", make_pdf_bytes("Surat keterangan"))
    doc_id = only_doc_id(sigs)
    signed = (docs / f"{doc_id}_signed.pdf").read_bytes()
    html = verify(c, signed, doc_id)
    assert "UTUH" in html
    assert STATUS_VALID in html


def test_one_byte_change_in_signed_pdf_fails(client):
    c, docs, sigs = client
    sign(c, OWNER_A, PASS_A, "Budi", make_pdf_bytes("Surat keterangan"))
    doc_id = only_doc_id(sigs)
    signed = bytearray((docs / f"{doc_id}_signed.pdf").read_bytes())
    signed[len(signed) // 2] ^= 0x01
    record = json.loads((sigs / f"{doc_id}.json").read_text())
    assert record["signed_hashes"]
    assert hashlib.sha256(bytes(signed)).hexdigest() not in record["signed_hashes"]
    html = verify(c, bytes(signed), doc_id)
    assert "BERUBAH" in html
    assert status_block("TAMPERED") in html


def test_multi_signer_keeps_all_versions_valid(client):
    c, docs, sigs = client
    sign(c, OWNER_A, PASS_A, "Budi", make_pdf_bytes("Lembar pengesahan"))
    doc_id = only_doc_id(sigs)
    version1 = (docs / f"{doc_id}_signed.pdf").read_bytes()
    sign(c, OWNER_B, PASS_B, "Siti", doc_id=doc_id)
    version2 = (docs / f"{doc_id}_signed.pdf").read_bytes()
    assert version1 != version2
    record = json.loads((sigs / f"{doc_id}.json").read_text())
    assert len(record["signers"]) == 2
    assert len(record["signed_hashes"]) == 2
    assert hashlib.sha256(version1).hexdigest() in app_module.known_hashes(record)
    assert hashlib.sha256(version2).hexdigest() in app_module.known_hashes(record)
    for version in (version1, version2):
        html = verify(c, version, doc_id)
        assert "UTUH" in html and STATUS_VALID in html


@pytest.mark.parametrize("extension", [".docx", ".jpg", ".jpeg", ".png", ".txt", ".xlsx"])
def test_non_pdf_sign_download_and_verify_roundtrip(client, extension):
    c, docs, sigs = client
    original = make_document_bytes(extension)
    sign(c, OWNER_A, PASS_A, "Budi", original, filename=f"surat{extension}")
    doc_id = only_doc_id(sigs)

    signed_path = docs / f"{doc_id}_signed{extension}"
    assert signed_path.is_file()
    downloaded = c.get(f"/download/{doc_id}")
    assert downloaded.status_code == 200
    assert downloaded.data == signed_path.read_bytes()
    assert downloaded.mimetype == app_module.document_module.MIME_TYPES[extension]
    assert c.get(f"/preview/{doc_id}").status_code == 404

    html = verify(c, downloaded.data, "", filename=f"hasil{extension}")
    assert "UTUH" in html
    assert STATUS_VALID in html

    tampered = tamper_document(downloaded.data, extension)
    tampered_html = verify(c, tampered, "", filename=f"diubah{extension}")
    assert "BERUBAH" in tampered_html
    assert status_block("TAMPERED") in tampered_html


def test_30_mib_text_upload_signs_and_verifies(client):
    c, docs, sigs = client
    payload = b"A" * app_module.MAX_SOURCE_DOCUMENT_SIZE
    response = sign(c, OWNER_A, PASS_A, "Budi", payload, filename="ukuran-tepat.txt")
    assert response.status_code == 200
    doc_id = only_doc_id(sigs)
    signed = c.get(f"/download/{doc_id}")
    assert signed.status_code == 200
    html = verify(c, signed.data, "", filename="ukuran-tepat_signed.txt")
    assert "UTUH" in html and STATUS_VALID in html
