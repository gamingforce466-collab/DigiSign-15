"""Validasi, penyisipan tanda tangan, dan ekstraksi QR untuk XLSX."""
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.styles import Font

from formats.common import extract_media_images, image_bytes, valid_ooxml_package
from qr import generator as qr_generator


def is_valid(data):
    if not valid_ooxml_package(data, "xl/workbook.xml"):
        return False
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, keep_links=False)
        workbook.close()
        return True
    except Exception:
        return False


def add_signatures(source_path, output_path, qr_images, signers):
    workbook = load_workbook(source_path, keep_links=True)
    sheet_name = "DigiSign"
    suffix = 2
    while sheet_name in workbook.sheetnames:
        sheet_name = f"DigiSign {suffix}"
        suffix += 1
    sheet = workbook.create_sheet(sheet_name)
    sheet["A1"] = "TANDA TANGAN DIGITAL"
    sheet["A1"].font = Font(name="Calibri", size=16, bold=True)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 36
    for index, (qr_image, signer) in enumerate(zip(qr_images, signers), start=1):
        row = 3 + (index - 1) * 12
        sheet.cell(row=row, column=1, value=f"Penandatangan #{index}").font = Font(name="Calibri", bold=True)
        for offset, (label, key) in enumerate((
            ("Nama", "name"),
            ("Jabatan", "position"),
            ("Institusi", "institution"),
            ("Tanggal", "date"),
        ), start=1):
            sheet.cell(row=row + offset, column=1, value=label)
            sheet.cell(row=row + offset, column=2, value=signer.get(key, ""))
        image = SpreadsheetImage(BytesIO(image_bytes(qr_image)))
        image.width = 220
        image.height = 220
        sheet.add_image(image, f"D{row}")
        sheet.row_dimensions[row].height = 24
    workbook.save(output_path)
    workbook.close()


def extract_metadata(path):
    return qr_generator.decode_qr_from_images(extract_media_images(path, "xl/media/"))